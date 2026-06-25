"""
Excel 파일에서 Event 데이터를 DB에 임포트합니다.

사용법:
  python manage.py import_events events.xlsx
  python manage.py import_events events.xlsx --clear   # 기존 데이터 먼저 삭제
"""

from datetime import date

from django.core.management.base import BaseCommand, CommandError

from contents.models import Event


class Command(BaseCommand):
    help = "Excel 파일에서 교육행사 데이터를 임포트합니다."

    def add_arguments(self, parser):
        parser.add_argument("file", help="임포트할 Excel 파일 경로 (events.xlsx)")
        parser.add_argument(
            "--clear",
            action="store_true",
            help="임포트 전 기존 Event 레코드를 모두 삭제합니다.",
        )

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl이 필요합니다: pip install openpyxl")

        filepath = options["file"]
        try:
            wb = openpyxl.load_workbook(filepath)
        except FileNotFoundError:
            raise CommandError(f"파일을 찾을 수 없습니다: {filepath}")

        ws = wb.active
        headers = [cell.value for cell in ws[1]]

        if options["clear"]:
            count, _ = Event.objects.all().delete()
            self.stdout.write(f"기존 {count}건 삭제 완료")

        VALID_STATUS = {"open", "closed", "ended"}
        VALID_ONLINE = {"offline", "online", "both"}

        created = skipped = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))

            title = str(data.get("title") or "").strip()
            if not title:
                skipped += 1
                continue

            def clean(field, default="", max_len=None):
                val = str(data.get(field) or default).strip()
                return val[:max_len] if max_len else val

            def parse_date_val(val):
                if val is None:
                    return None
                if isinstance(val, date):
                    return val
                s = str(val).strip()
                if not s:
                    return None
                try:
                    return date.fromisoformat(s[:10])
                except ValueError:
                    return None

            status = clean("status", "open")
            if status not in VALID_STATUS:
                status = "open"

            online_type = clean("online_type", "offline")
            if online_type not in VALID_ONLINE:
                online_type = "offline"

            Event.objects.create(
                title=title[:200],
                summary=clean("summary", max_len=300),
                body=clean("body"),
                status=status,
                region=clean("region", max_len=30),
                online_type=online_type,
                host=clean("host", max_len=100),
                start_date=parse_date_val(data.get("start_date")),
                end_date=parse_date_val(data.get("end_date")),
                place_name=clean("place_name", max_len=120),
                address=clean("address", max_len=255),
            )
            created += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"임포트 완료: {created}건 생성, {skipped}건 스킵"
            )
        )
