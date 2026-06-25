"""
Data/lifecycle_edu_youth.xlsx 에서 청년기 경제교육 정보를 Content 모델에 임포트합니다.

사용법:
  python manage.py import_lifecycle_edu
  python manage.py import_lifecycle_edu --file Data/lifecycle_edu_youth.xlsx
  python manage.py import_lifecycle_edu --clear   # 기존 임포트 데이터 삭제 후 재적재
"""

from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError

from contents.models import Content

SOURCE_TAG = "[생애주기경제교육]"


class Command(BaseCommand):
    help = "생애주기별 경제교육 Excel 데이터를 Content 모델에 임포트합니다."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default="Data/lifecycle_edu_youth.xlsx",
            help="임포트할 Excel 파일 경로 (기본값: Data/lifecycle_edu_youth.xlsx)",
        )
        parser.add_argument(
            "--clear",
            action="store_true",
            help="summary 에 SOURCE_TAG 가 포함된 기존 레코드를 삭제한 뒤 임포트",
        )

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl이 필요합니다: pip install openpyxl")

        filepath = Path(settings.BASE_DIR) / options["file"]
        if not filepath.exists():
            raise CommandError(
                f"파일을 찾을 수 없습니다: {filepath}\n"
                "먼저 python manage.py fetch_lifecycle_edu 를 실행하세요."
            )

        if options["clear"]:
            deleted, _ = Content.objects.filter(summary__contains=SOURCE_TAG).delete()
            self.stdout.write(f"기존 {deleted}건 삭제 완료")

        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]

        created = skipped = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))

            institution = str(data.get("기관명") or "").strip()
            if not institution:
                skipped += 1
                continue

            offline = str(data.get("오프라인 교육") or "").strip()
            online = str(data.get("온라인 교육") or "").strip()
            link = str(data.get("링크주소") or "").strip()
            stage = str(data.get("생애주기") or "청년기").strip()

            title = f"{institution} {stage} 경제교육"

            body_parts = []
            if offline:
                body_parts.append(f"[오프라인 교육]\n{offline}")
            if online:
                body_parts.append(f"[온라인 교육]\n{online}")
            body = "\n\n".join(body_parts)

            summary = f"{SOURCE_TAG} {stage} | {institution}"

            Content.objects.create(
                title=title[:200],
                summary=summary[:300],
                body=body,
                category="economy",
                external_url=link,
                is_recommended=True,
            )
            self.stdout.write(f"  + {title}")
            created += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"\n임포트 완료: {created}건 생성, {skipped}건 스킵"
            )
        )
